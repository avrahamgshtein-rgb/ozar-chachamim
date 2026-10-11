#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Add one research paper (already cleaned to plain lines) to a sage that is on the roster.

    python3 scripts/ingest_drive_doc.py --text paper.txt --sage-id 622 \
        --sage-label "הרב יהודה לייב הלוי אשלג (בעל הסולם)"

The text file is the paper as plain lines: line 1 is the title, every other line
is one paragraph (a table row is "cell | cell"). Markdown debris and the Arabic
letters that Google's export sometimes leaves inside Hebrew words must be gone
before this runs; the script refuses a file that still has them.

It does, in this order, everything a paper needs and nothing else:
  1. saves the paper as data/<title>.docx, the source the corpus is rebuilt from
  2. records it in data/research_attribution.json, keyed by the sage's NAME
  3. appends it to nextjs-app/public/research/<id>.json (never replaces)
  4. sets has_research on the sage in both copies of data.json

Why it appends instead of running rebuild_research_corpus.py: that rebuild does
not reproduce the files now in the repo (25 differ, because later fixes were
made to the JSON by hand), so running it would undo those fixes.

The sage is checked by id AND label, because ids are what drift between data
sets and names are what stay put. Re-running with the same paper is a no-op.
"""
import argparse, io, json, os, re, sys
from docx import Document

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_FILES = [os.path.join(REPO, 'data.json'), os.path.join(REPO, 'nextjs-app', 'public', 'data.json')]
RESEARCH = os.path.join(REPO, 'nextjs-app', 'public', 'research')
ATTR = os.path.join(REPO, 'data', 'research_attribution.json')

ap = argparse.ArgumentParser()
ap.add_argument('--text', required=True)
ap.add_argument('--sage-id', required=True)
ap.add_argument('--sage-label', required=True)
ap.add_argument('--docx-name', help='file name under data/; defaults to one made from the title')
ap.add_argument('--dry-run', action='store_true')
args = ap.parse_args()

def die(msg):
    print('ERROR:', msg, file=sys.stderr)
    sys.exit(1)

def read(path):
    return io.open(path, encoding='utf-8').read()

lines = [l.strip() for l in read(args.text).split('\n') if l.strip()]
if len(lines) < 5:
    die('the text has fewer than 5 non-empty lines')
title = lines[0]
content = '\n'.join(lines)

arabic = re.findall(r'[؀-ۿݐ-ݿ]+', content)
if arabic:
    die(f'Arabic letters inside the text: {arabic[:5]}')
if re.search(r'\*\*|^#|\\\[|\\"', content, flags=re.M):
    die('markdown debris left in the text (**, #, \\[ or \\")')

docx_name = args.docx_name or (re.sub(r'[\\/*?<>|"]', '', title.replace(':', '-'))[:180].strip() + '.docx')

# the sage must be on the roster under this id with this label
datasets = [json.loads(read(p)) for p in DATA_FILES]
for p, d in zip(DATA_FILES, datasets):
    node = next((n for n in d['nodes'] if n['id'] == args.sage_id), None)
    if not node:
        die(f'no sage with id {args.sage_id} in {p}')
    if node['label'] != args.sage_label:
        die(f'id {args.sage_id} is "{node["label"]}" in {p}, not "{args.sage_label}"')

research_path = os.path.join(RESEARCH, f'{args.sage_id}.json')
existing = json.loads(read(research_path)) if os.path.exists(research_path) else []
if any(e.get('source_file') == docx_name or e.get('title') == title for e in existing):
    print(f'already ingested for {args.sage_id}: {title[:60]}')
    sys.exit(0)

attr = json.loads(read(ATTR))
entry_exists = any(e.get('file') == docx_name for e in attr['documents'])

if args.dry_run:
    print(f'would add {len(content.split()):,} words to sage {args.sage_id} as {docx_name}')
    sys.exit(0)

# 1. the .docx
docx_path = os.path.join(REPO, 'data', docx_name)
doc = Document()
for l in lines:
    doc.add_paragraph(l)
doc.save(docx_path)
back = '\n'.join(p.text.strip() for p in Document(docx_path).paragraphs if p.text.strip())
if back != content:
    os.remove(docx_path)
    die('the .docx does not read back as the same text')

# 2. attribution, in the file's own formatting (one-space indent)
if not entry_exists:
    attr['documents'].append({'file': docx_name, 'sage_label': args.sage_label, 'confidence': 'high', 'evidence': title})
    before = read(ATTR)
    if json.dumps(json.loads(before), ensure_ascii=False, indent=1) != before.rstrip('\n'):
        die('research_attribution.json does not round-trip, so it would be reformatted; edit it by hand')
    io.open(ATTR, 'w', encoding='utf-8').write(json.dumps(attr, ensure_ascii=False, indent=1) + ('\n' if before.endswith('\n') else ''))

# 3. the research file, appended
if existing and json.dumps(existing, ensure_ascii=False, indent=2) != read(research_path):
    die(f'{research_path} does not round-trip, so appending would reformat it')
existing.append({'title': title[:200], 'source_file': docx_name, 'word_count': len(content.split()), 'content': content})
io.open(research_path, 'w', encoding='utf-8').write(json.dumps(existing, ensure_ascii=False, indent=2))

# 4. has_research on both copies of data.json
for p, d in zip(DATA_FILES, datasets):
    raw = read(p)
    if json.dumps(json.loads(raw), ensure_ascii=False, indent=1) != raw.rstrip('\n'):
        die(f'{p} does not round-trip, so it would be reformatted')
    for n in d['nodes']:
        if n['id'] == args.sage_id:
            n['has_research'] = True
    io.open(p, 'w', encoding='utf-8').write(json.dumps(d, ensure_ascii=False, indent=1) + ('\n' if raw.endswith('\n') else ''))

print(f'ingested {len(content.split()):,} words for {args.sage_id} ({args.sage_label}): {docx_name}')
