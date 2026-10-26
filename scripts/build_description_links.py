#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Builds the file the site owner works through to put a link to the site into the description of every
YouTube video and every Spotify episode: data/description-links.xlsx (what he opens) and its sidecar
data/description-links.json (what tests/descriptionLinks.test.ts reads).

    python3 scripts/build_description_links.py             # fetch the views, write both files
    python3 scripts/build_description_links.py --no-stats  # no network: the YouTube sheet stays in playlist order
    python3 scripts/build_description_links.py --check     # rebuild in memory, exit 1 when a committed file differs
    python3 scripts/build_description_links.py --site https://example.org --out DIR

YouTube and Spotify descriptions cannot be edited with the read-only API key this repo has, so the
deliverable is a list: one row per mapped video (sorted by views, most viewed first, so the first rows
bring the most readers) and one per podcast episode, each with the edit link and ONE line to paste at the
top of the description. The line links to that sage's page in the language of the video:

    he  📜 רש״י (רבי שלמה יצחקי) באוצר חכמים — המחקר המלא, חידון ורבותיו ותלמידיו: <url>
    en  📜 Rashi on Ozar Chachamim — the full research, a quiz, and his teachers and students: <url>
    ru  📜 Раши в «Оцар Хахамим» — полное исследование, викторина, а также его учителя и ученики: <url>

The url is <site>/<lang>/sage/<id>?utm_source=youtube&utm_medium=video (Spotify: ...spotify&...podcast, the
podcast is Hebrew), so Vercel Analytics shows what each channel brings. <site> is DEFAULT_SITE_URL of
nextjs-app/lib/siteUrl.ts (one source of truth), or $NEXT_PUBLIC_SITE_URL when that is set, or --site.

Nothing in a line is a guess. Each optional part is there only when it is true for that sage in that language:
  the full research   public/research/<id>.json (he) / <id>.<lang>.json exists and holds text; otherwise
                      "his page" (a woman: "her page"; a book or a topic: "the page"),
  a quiz              public/data-quiz.json lists the sage for that language,
  teachers, students  the sage has a teacher / a student link in the corpus the app itself serves
                      (lib/serverData.ts::getAllConnections(), which merges data.json and the supplement and
                      link files exactly as the pages do; read through a small `tsx` run). Only the half that is
                      true is said: "his teachers", "his students" or both. An entry that is a book or a topic
                      (chapter_type in NOT_A_PERSON, as in rebuild_data_from_csv.py) never claims either.
Which video is which sage is NOT guessed: data/youtube-sage-map.json (hand-checked, by playlist row) says, and
the video's language is its own `defaultLanguage` from data/youtube-playlist.tsv (never inferred from the
title: see scripts/snapshot_youtube.py). Rows in `_left_out` are not listed. The Hebrew name is the roster
label without its subtitle (lib/displayName.ts), with proper gershayim; en/ru use the overlay label.

Grammatical gender: a woman is a sage whose tags (as the app serves them) hold 'נשים', and nothing else
decides it. (Until 2026-10-25 five women lacked the tag and a list of their names stood in for it here; they
were tagged in the roster, so the list is gone.)

Views come from the YouTube Data API (videos.list, 50 ids a call), called the way scripts/snapshot_youtube.py
does: curl, no key unless $YOUTUBE_API_KEY is set (the cloud environment's network secret adds the key as an
X-Goog-Api-Key header). When the API fails the files are still written, in playlist order, and the script says so.
--check does not need the network: it rebuilds every row and line and compares them with the committed JSON and
xlsx, ignoring the views (they change daily) and the date they were read, and checks the sort order is by views.
Never hand-edit the outputs: regenerate.
"""
import argparse, datetime, json, os, re, shutil, subprocess, sys

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP = os.path.join(REPO, 'nextjs-app')
PUBLIC = os.path.join(APP, 'public')
TSV = os.path.join(REPO, 'data', 'youtube-playlist.tsv')
MAP = os.path.join(REPO, 'data', 'youtube-sage-map.json')
ROSTER = os.path.join(PUBLIC, 'data.json')
QUIZ_INDEX = os.path.join(PUBLIC, 'data-quiz.json')
SITE_URL_TS = os.path.join(APP, 'lib', 'siteUrl.ts')
API = 'https://youtube.googleapis.com/youtube/v3'
LANGS = ('he', 'en', 'ru')
API_LANG = {'ru': 'ru', 'en': 'en', 'iw': 'he', 'he': 'he'}
LINE_LIMIT = 200          # a line is shorter than this, the url included, counted in UTF-16 units (as JavaScript's .length does: the emoji is two)
NOT_A_PERSON = {'נושא/מושג', 'נושא', 'ספר', 'חיבור', 'אירוע היסטורי', 'דמויות'}   # as in rebuild_data_from_csv.py
TEXTS = {'מחזור ויטרי', 'מדרש רבה'}   # typed 'טקסט יסוד' like Rav Sherira Gaon and Ben Sira, but books (as in build_catalog_links.py)

ap = argparse.ArgumentParser()
ap.add_argument('--check', action='store_true', help='rebuild in memory and exit 1 when the committed files differ')
ap.add_argument('--no-stats', action='store_true', help='do not ask YouTube for the views (playlist order)')
ap.add_argument('--site', help='the site address in the lines (default: $NEXT_PUBLIC_SITE_URL, else DEFAULT_SITE_URL of nextjs-app/lib/siteUrl.ts)')
ap.add_argument('--out', default=os.path.join(REPO, 'data'), help='directory of description-links.xlsx / .json')
args = ap.parse_args()
OUT_XLSX = os.path.join(args.out, 'description-links.xlsx')
OUT_JSON = os.path.join(args.out, 'description-links.json')

# ---- the wording --------------------------------------------------------------------------------------------
# kind: 'm' / 'f' a man / a woman, 'n' a book, a group or a topic (no possessive, never a teachers/students claim)
SITE_NAME = {'he': 'אוצר חכמים', 'en': 'Ozar Chachamim', 'ru': '«Оцар Хахамим»'}
HEAD = {   # {name} {where}
    'he': '📜 {name} באוצר חכמים — {parts}: {url}',
    'en': '📜 {name} on Ozar Chachamim — {parts}: {url}',
    'ru': '📜 {name} в «Оцар Хахамим» — {parts}: {url}',
}
RESEARCH = {'he': 'המחקר המלא', 'en': 'the full research', 'ru': 'полное исследование'}
PAGE = {
    'he': {'m': 'הדף שלו', 'f': 'הדף שלה', 'n': 'הדף'},
    'en': {'m': 'his page', 'f': 'her page', 'n': 'the page'},
    'ru': {'m': 'его страница', 'f': 'её страница', 'n': 'страница'},
}
QUIZ = {'he': 'חידון', 'en': 'a quiz', 'ru': 'викторина'}
LINKS = {   # (teachers, students) -> text
    'he': {'m': {(1, 1): 'רבותיו ותלמידיו', (1, 0): 'רבותיו', (0, 1): 'תלמידיו'},
           'f': {(1, 1): 'רבותיה ותלמידיה', (1, 0): 'רבותיה', (0, 1): 'תלמידיה'}},
    'en': {'m': {(1, 1): 'his teachers and students', (1, 0): 'his teachers', (0, 1): 'his students'},
           'f': {(1, 1): 'her teachers and students', (1, 0): 'her teachers', (0, 1): 'her students'}},
    'ru': {'m': {(1, 1): 'его учителя и ученики', (1, 0): 'его учителя', (0, 1): 'его ученики'},
           'f': {(1, 1): 'её учителя и ученики', (1, 0): 'её учителя', (0, 1): 'её ученики'}},
}


def join_parts(lang, items):
    """'A', 'A and B', 'A, B and C' in the language's way; the links part (always last) is a phrase with its own 'and'."""
    if len(items) == 1:
        return items[0]
    head, last = items[:-1], items[-1]
    if lang == 'he':
        return ', '.join(head) + ' ו' + last
    if lang == 'en':
        return ', '.join(head) + (', and ' if len(items) > 2 else ' and ') + last
    # Russian: "A, B, а также её учителя и ученики" when the last part holds its own "и"; else "A и B"
    if ' и ' in last:
        return ', '.join(head) + ', а также ' + last
    return ', '.join(head) + ' и ' + last


def gershayim(text):
    """A quote between two Hebrew letters is a gershayim, one after a Hebrew letter a geresh (lib/pathShare.ts::gershayim)."""
    return re.sub(r"(?<=[א-ת])'", '׳', re.sub(r'(?<=[א-ת])"(?=[א-ת])', '״', text))


def make_line(lang, name, url, kind, research, quiz, teachers, students):
    items = [RESEARCH[lang] if research else PAGE[lang][kind]]
    if quiz:
        items.append(QUIZ[lang])
    if kind != 'n' and (teachers or students):
        items.append(LINKS[lang][kind][(int(teachers), int(students))])
    return HEAD[lang].format(name=name, parts=join_parts(lang, items), url=url)


def utf16_len(text):
    return len(text.encode('utf-16-le')) // 2


def short_name(name):
    """The name without a trailing parenthetical alias, for the rare line that would not fit."""
    return re.sub(r'\s*\([^)]*\)\s*$', '', name).strip() or name


# ---- inputs -------------------------------------------------------------------------------------------------
def read_site():
    """--site, else $NEXT_PUBLIC_SITE_URL (the app's one switch, normalised as lib/siteUrl.ts does), else DEFAULT_SITE_URL."""
    m = re.search(r"const DEFAULT_SITE_URL = '([^']+)'", open(SITE_URL_TS, encoding='utf-8').read())
    if not m:
        sys.exit('cannot read DEFAULT_SITE_URL from nextjs-app/lib/siteUrl.ts')
    raw = (args.site or os.environ.get('NEXT_PUBLIC_SITE_URL') or m.group(1)).strip().rstrip('/')
    if not raw:
        raw = m.group(1)
    return raw if re.match(r'https?://', raw) else 'https://' + raw


def read_playlist():
    rows = {}
    for line in open(TSV, encoding='utf-8').read().splitlines():
        n, vid, sec, lang, title = line.split('\t')
        if lang not in LANGS:
            sys.exit('playlist row %s has the language "%s": refresh the snapshot (scripts/snapshot_youtube.py)' % (n, lang))
        rows[int(n)] = {'id': vid, 'lang': lang, 'title': title}
    return rows


def left_out_rows(left_out):
    out = set()
    for key in left_out:
        for part in key.split(','):
            ends = part.split('-')
            out |= set(range(int(ends[0]), int(ends[-1]) + 1))
    return out


# What the app itself serves. A `tsx` run over the app's own modules, so that "the sage has a teacher / a student" is
# computed from the merged corpus (data.json + supplements + the generated link files + the patch), with the same
# merge the pages use, and so that the names are the ones the pages show (localizeSage + labelParts).
FACTS_TS = r'''
import { getAllSages, getAllConnections } from '__APP__/lib/serverData'
import { localizeSage } from '__APP__/lib/serverDataTranslator'
import { labelParts } from '__APP__/lib/displayName'
const sages = getAllSages()
const known = new Set(sages.map(s => s.id))
const hasTeacher = new Set<string>()
const hasStudent = new Set<string>()
for (const l of getAllConnections()) {
  if (l.source === l.target || !known.has(l.source) || !known.has(l.target)) continue
  if (l.type === 'teacher') { hasStudent.add(l.source); hasTeacher.add(l.target) }
  else if (l.type === 'student') { hasTeacher.add(l.source); hasStudent.add(l.target) }
}
const out: Record<string, unknown> = {}
for (const s of sages) {
  out[s.id] = {
    label: s.label,
    tags: s.tags ?? [],
    names: {
      he: labelParts(s.label).name,
      en: labelParts(localizeSage(s, 'en').label).name,
      ru: labelParts(localizeSage(s, 'ru').label).name,
    },
    teachers: hasTeacher.has(s.id),
    students: hasStudent.has(s.id),
  }
}
console.log(JSON.stringify(out))
'''


def app_facts():
    local = os.path.join(APP, 'node_modules', '.bin', 'tsx')
    cmd = [local] if os.path.exists(local) else [shutil.which('npx') or 'npx', '--yes', 'tsx']
    code = FACTS_TS.replace('__APP__', APP)
    r = subprocess.run(cmd + ['-e', code], cwd=APP, capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit('could not read the corpus through tsx:\n' + r.stderr[-2000:])
    return json.loads(r.stdout)


def research_ok(sid, lang):
    p = os.path.join(PUBLIC, 'research', '%s.json' % sid if lang == 'he' else '%s.%s.json' % (sid, lang))
    if not os.path.exists(p):
        return False
    docs = json.load(open(p, encoding='utf-8'))
    return isinstance(docs, list) and any(str(d.get('content', '')).strip() for d in docs)


def build_doc(site):
    rows = read_playlist()
    mp = json.load(open(MAP, encoding='utf-8'))
    roster = {n['id']: n for n in json.load(open(ROSTER, encoding='utf-8'))['nodes']}
    quiz = {sid: set(v['langs']) for sid, v in json.load(open(QUIZ_INDEX, encoding='utf-8'))['sages'].items()}
    facts = app_facts()

    for sid, node in roster.items():       # the app serves the roster the script reads
        if sid not in facts or facts[sid]['label'] != node['label']:
            sys.exit('sage %s: the app says "%s", data.json says "%s"' % (sid, facts.get(sid, {}).get('label'), node['label']))

    def kind_of(sid):
        node = roster[sid]
        if node.get('chapter_type') in NOT_A_PERSON or node['label'] in TEXTS:
            return 'n'
        return 'f' if 'נשים' in facts[sid]['tags'] else 'm'

    def claims_of(sid, lang):
        f = facts[sid]
        kind = kind_of(sid)
        return {'kind': kind, 'research': research_ok(sid, lang), 'quiz': lang in quiz.get(sid, ()),
                'teachers': bool(f['teachers']) and kind != 'n', 'students': bool(f['students']) and kind != 'n'}

    def line_of(sid, lang, src, medium):
        url = '%s/%s/sage/%s?utm_source=%s&utm_medium=%s' % (site, lang, sid, src, medium)
        c = claims_of(sid, lang)
        name = facts[sid]['names'][lang]
        if lang == 'he':
            name = gershayim(name)
        for candidate in (name, short_name(name)):
            line = make_line(lang, candidate, url, c['kind'], c['research'], c['quiz'], c['teachers'], c['students'])
            if utf16_len(line) < LINE_LIMIT:
                return candidate, line, c
        sys.exit('a line of %d characters (limit %d) for sage %s in %s: %s' % (utf16_len(line), LINE_LIMIT, sid, lang, line))

    # ---- YouTube: one entry per mapped row ----
    owner = {}
    for sid, rs in mp['sages'].items():
        if sid not in roster:
            sys.exit('the map names sage %s, who is not on the roster' % sid)
        for r in rs:
            if r in owner:
                sys.exit('playlist row %d is mapped to sages %s and %s' % (r, owner[r], sid))
            if r not in rows:
                sys.exit('the map names playlist row %d, which is not in the snapshot' % r)
            owner[r] = sid
    skipped = left_out_rows(mp['_left_out'])
    clash = sorted(set(owner) & skipped)
    if clash:
        sys.exit('rows %s are both mapped and left out' % clash)
    loose = sorted(set(rows) - set(owner) - skipped)
    if loose:
        sys.exit('playlist rows %s are neither mapped nor left out: update data/youtube-sage-map.json' % loose)

    youtube = []
    for r in sorted(owner):
        sid, v = owner[r], rows[r]
        name, line, c = line_of(sid, v['lang'], 'youtube', 'video')
        youtube.append({'row': r, 'id': v['id'], 'lang': v['lang'], 'sage': sid, 'label': roster[sid]['label'],
                        'name': name, 'title': v['title'], 'claims': c, 'line': line})

    # ---- Spotify: a sage whose spotify_url is an episode address (the podcast is Hebrew) ----
    spotify, unresolved = [], []
    for sid in sorted(roster, key=int):
        url = (roster[sid].get('spotify_url') or '').strip()
        m = re.match(r'https://open\.spotify\.com/episode/([A-Za-z0-9]+)', url)
        if m:
            name, line, c = line_of(sid, 'he', 'spotify', 'podcast')
            spotify.append({'sage': sid, 'label': roster[sid]['label'], 'name': name,
                            'episode': 'https://open.spotify.com/episode/' + m.group(1), 'claims': c, 'line': line})
        elif url:
            unresolved.append({'sage': sid, 'label': roster[sid]['label'], 'link': url})

    return {
        '_about': 'Generated by scripts/build_description_links.py from data/youtube-playlist.tsv, data/youtube-sage-map.json, '
                  'the roster (nextjs-app/public/data.json) and the corpus the app serves. Do not edit: regenerate. '
                  'One line per mapped video (in the video\'s own language) and one per Spotify episode, to paste at the top '
                  'of the description; the xlsx beside it is what the owner works through.',
        'site': site,
        'youtube': youtube,
        'spotify': spotify,
        'unresolved': unresolved,
    }


# ---- views --------------------------------------------------------------------------------------------------
def api_get(path, **params):
    import urllib.parse
    if os.environ.get('YOUTUBE_API_KEY'):
        params['key'] = os.environ['YOUTUBE_API_KEY']
    url = API + path + '?' + urllib.parse.urlencode(params)
    last = ''
    for attempt in range(3):
        r = subprocess.run(['curl', '-s', '-m', '60', url], capture_output=True, text=True)
        try:
            d = json.loads(r.stdout)
        except ValueError:
            last = 'not JSON: %s' % (r.stdout[:120] or r.stderr[:120] or 'empty answer')
            continue
        if 'error' in d:
            last = d['error'].get('message', str(d['error']))
            if d['error'].get('code') in (400, 403, 404):
                break
            continue
        return d
    raise RuntimeError(last)


def fetch_views(doc):
    """{video id: views or None} from the API, plus warnings; (None, reason) when the API cannot be used."""
    ids = [e['id'] for e in doc['youtube']]
    views, warnings = {}, []
    try:
        for i in range(0, len(ids), 50):
            d = api_get('/videos', part='statistics,snippet', maxResults=50, id=','.join(ids[i:i + 50]))
            for it in d['items']:
                views[it['id']] = it
    except Exception as e:     # network, quota, key: the file is still written, in playlist order
        return None, str(e)
    out = {}
    for e in doc['youtube']:
        it = views.get(e['id'])
        if it is None:
            warnings.append('row %d (%s): YouTube does not return the video (deleted or private?)' % (e['row'], e['id']))
            out[e['id']] = None
            continue
        count = it.get('statistics', {}).get('viewCount')
        out[e['id']] = int(count) if count is not None else None
        own = API_LANG.get(it.get('snippet', {}).get('defaultLanguage'), '')
        if own and own != e['lang']:
            warnings.append('row %d (%s): the snapshot says %s, YouTube says %s' % (e['row'], e['id'], e['lang'], own))
    return out, warnings


# ---- the workbook -------------------------------------------------------------------------------------------
LANG_NAME = {'he': 'עברית', 'en': 'English', 'ru': 'Русский'}
YT_HEAD = ['#', 'צפיות', 'שפה', 'החכם', 'כותרת הסרטון', 'עריכה ב־YouTube Studio', 'השורה להדבקה', 'הסרטון', 'שורה בפלייליסט', 'הודבק ✓']
SP_HEAD = ['#', 'החכם', 'הפרק בספוטיפיי', 'השורה להדבקה', 'צפיות ביוטיוב בסך הכול (לסדר הרשימה)', 'מזהה החכם', 'הודבק ✓']
VIEWS_NOTE_AT = 2          # the row of the instructions that dates the views: --check does not compare it
UNRESOLVED_TITLE = 'קישורים שלא נפתחים'
UNRESOLVED_NOTE = ('הקישור שבקובץ הוא קישור של Spotify for Creators ולא כתובת של פרק, ולכן אין עבורו שורה. '
                   'כשתגיע הכתובת של הפרק, החכם יתווסף לגיליון.')


def studio_link(vid):
    return 'https://studio.youtube.com/video/%s/edit' % vid


def video_link(vid):
    return 'https://youtu.be/%s' % vid


def yt_cells(e):
    """The cells of a YouTube row that do not depend on the views (the same list writes and checks)."""
    return {3: LANG_NAME[e['lang']], 4: e['label'], 5: e['title'], 6: studio_link(e['id']), 7: e['line'], 8: video_link(e['id']), 9: e['row']}


def sp_cells(e):
    return {2: e['label'], 3: e['episode'], 4: e['line'], 6: int(e['sage'])}


def order_youtube(doc, views):
    key = (lambda e: (-(views.get(e['id']) or -1), e['row'])) if views else (lambda e: e['row'])
    return sorted(doc['youtube'], key=key)


def order_spotify(doc, views):
    total = spotify_totals(doc, views)
    return sorted(doc['spotify'], key=lambda e: (-total.get(e['sage'], 0), int(e['sage'])))


def spotify_totals(doc, views):
    total = {}
    if views:
        for e in doc['youtube']:
            total[e['sage']] = total.get(e['sage'], 0) + (views.get(e['id']) or 0)
    return total


def instructions(doc, views_note):
    """Rows of the instructions sheet: (text, link or None, style). views_note is the one row --check ignores."""
    site = doc['site']
    u = lambda path, src, med, extra='': '%s%s?utm_source=%s&utm_medium=%s%s' % (site, path, src, med, extra)
    n_sp, n_un = len(doc['spotify']), len(doc['unresolved'])
    return [
        ('איך משתמשים בקובץ', None, 'title'),
        ('המטרה: להוסיף לכל סרטון וכל פרק שורה אחת עם קישור לדף החכם באתר "אוצר חכמים", בשפת הסרטון, כדי שמי שכבר צופה ושומע יגיע אל המחקר המלא. '
         'הקישורים נושאים תגית (utm), ולכן ב־Vercel Analytics רואים כמה מגיעים מיוטיוב וכמה מספוטיפיי.', None, 'text'),
        (views_note, None, 'note'),
        ('', None, 'text'),
        ('יוטיוב (הגיליון "יוטיוב")', None, 'head'),
        ('1. הרשימה ממוינת לפי מספר הצפיות, מהרב־צפיות אל הפחות. מתחילים מלמעלה: השורות הראשונות מביאות הכי הרבה קוראים.', None, 'text'),
        ('2. לוחצים על הקישור בעמודה "עריכה ב־YouTube Studio" (צריך להיות מחוברים לערוץ). נפתח עמוד העריכה של הסרטון.', None, 'text'),
        ('3. מעתיקים את התא "השורה להדבקה", ומדביקים אותו בתחילת התיאור, מעל כל טקסט אחר (רק ההתחלה נראית לפני "עוד"), '
         'ושורה ריקה אחריו. לוחצים "שמירה".', None, 'text'),
        ('4. מסמנים ✓ בעמודה האחרונה, כדי לזכור איפה עצרתם.', None, 'text'),
        ('5. אין לשנות את הכתובת: התגית שבסופה היא מה שמאפשר לראות מאיפה הגיעו. סרטון ברוסית מקשר לדף ברוסית, סרטון באנגלית לדף באנגלית.', None, 'text'),
        ('', None, 'text'),
        ('ספוטיפיי (הגיליון "ספוטיפיי")', None, 'head'),
        ('1. שורה לכל פרק של הפודקאסט (%d פרקים), בעברית. הקישור בעמודה "הפרק בספוטיפיי" מראה איזה פרק זה.' % n_sp, None, 'text'),
        ('2. ב־Spotify for Creators: הפודקאסט, "פרקים", בוחרים את הפרק, עורכים את הפרטים, ומדביקים את השורה בתחילת התיאור. שומרים.', None, 'text'),
        ('3. הסדר הוא לפי סך הצפיות של אותו חכם ביוטיוב, כקירוב למי שהקהל רוצה לשמוע עליו (אין נתוני האזנות ב־API).', None, 'text'),
        ('4. %d קישורים שאינם כתובת פרק מופיעים בתחתית הגיליון, בלי שורה.' % n_un, None, 'text'),
        ('', None, 'text'),
        ('קישורים לכל הערוץ ולכל התוכנית (פעם אחת)', None, 'head'),
        ('ביוטיוב: בהתאמה אישית של הערוץ, בלשונית "פרופיל", בקטע "קישורים", להוסיף שלושה קישורים לדף הבית (עברית, English, Русский):', None, 'text'),
        ('אוצר חכמים', u('/he', 'youtube', 'channel'), 'link'),
        ('Ozar Chachamim', u('/en', 'youtube', 'channel'), 'link'),
        ('Оцар Хахамим', u('/ru', 'youtube', 'channel'), 'link'),
        ('ביוטיוב, בהגדרות הערוץ ("העלאת סרטונים", "ברירות מחדל"), אפשר להוסיף את אותו קישור לתיאור ברירת המחדל, כדי שסרטונים חדשים יישאו אותו מההתחלה.', None, 'text'),
        ('בספוטיפיי: בתיאור התוכנית (Show description) להוסיף קישור לדף הבית בעברית:', None, 'text'),
        ('אוצר חכמים', u('/he', 'spotify', 'podcast', '&utm_content=show'), 'link'),
        ('', None, 'text'),
        ('מה כתוב בשורות, ומה לא', None, 'head'),
        ('"המחקר המלא", "חידון" ו"רבותיו ותלמידיו" מופיעים בשורה רק כשהם קיימים בדף של אותו חכם ובאותה שפה. כשאין מחקר בשפה, כתוב "הדף שלו". '
         'לכן שורות של סרטונים שונים אינן זהות, וכך צריך להשאיר אותן.', None, 'text'),
        ('', None, 'text'),
        ('לתחזוקה: אם נוספו סרטונים או חכמים, מריצים מחדש python3 scripts/build_description_links.py (בבדיקה: --check). אסור לערוך את הקבצים ביד.', None, 'note'),
    ]


def views_note_text(date, ok):
    if ok:
        return 'נתוני הצפיות נלקחו מ־YouTube ב־%s.' % date
    return 'הקובץ נבנה בלי נתוני צפיות: סדר הגיליון "יוטיוב" הוא סדר הסרטונים ברשימת ההשמעה.'


def write_xlsx(doc, views, views_date):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font, head_fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='5B4B8A')
    wb = Workbook()

    def fit_to_page(ws):     # printed (or saved as PDF): landscape, one page wide
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def sheet(ws, headers, widths):
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = 'A2'
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        fit_to_page(ws)
        ws.print_title_rows = '1:1'

    def line_cell(c, lang):
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right' if lang == 'he' else 'left',
                                readingOrder=2 if lang == 'he' else 1)

    def link_cell(c, url):
        c.value, c.hyperlink = url, url
        c.font = Font(color='0563C1', underline='single')
        c.alignment = Alignment(horizontal='left', vertical='top', readingOrder=1)

    # 1. YouTube
    ws = wb.active
    ws.title = 'יוטיוב'
    sheet(ws, YT_HEAD, [5, 9, 9, 28, 44, 52, 100, 31, 10, 9])
    for rank, e in enumerate(order_youtube(doc, views), 1):
        cells = yt_cells(e)
        ws.append([rank, (views or {}).get(e['id'])] + [cells[i] for i in range(3, 10)] + [None])
        r = ws.max_row
        for col in (4, 5, 6, 7, 8):
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical='top')
        for col in (1, 2, 3, 9, 10):
            ws.cell(r, col).alignment = Alignment(horizontal='center', vertical='top')
        link_cell(ws.cell(r, 6), cells[6])
        link_cell(ws.cell(r, 8), cells[8])
        line_cell(ws.cell(r, 7), e['lang'])
        ws.cell(r, 5).alignment = Alignment(wrap_text=True, vertical='top', horizontal='left' if e['lang'] != 'he' else 'right',
                                            readingOrder=2 if e['lang'] == 'he' else 1)
        ws.cell(r, 2).number_format = '#,##0'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(YT_HEAD)), ws.max_row)

    # 2. Spotify
    ws = wb.create_sheet('ספוטיפיי')
    sheet(ws, SP_HEAD, [5, 34, 44, 100, 18, 9, 9])
    totals = spotify_totals(doc, views)
    for rank, e in enumerate(order_spotify(doc, views), 1):
        cells = sp_cells(e)
        ws.append([rank, cells[2], cells[3], cells[4], totals.get(e['sage'], 0) if views else None, cells[6], None])
        r = ws.max_row
        link_cell(ws.cell(r, 3), cells[3])
        line_cell(ws.cell(r, 4), 'he')
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
        for col in (1, 5, 6, 7):
            ws.cell(r, col).alignment = Alignment(horizontal='center', vertical='top')
        ws.cell(r, 5).number_format = '#,##0'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(SP_HEAD)), ws.max_row)
    if doc['unresolved']:
        ws.append([])
        ws.append([None, UNRESOLVED_TITLE])
        ws.cell(ws.max_row, 2).font = Font(bold=True, size=12)
        ws.append([None, UNRESOLVED_NOTE])
        ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=4)
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[ws.max_row].height = 32
        for e in doc['unresolved']:
            ws.append([None, e['label'], e['link'], None, None, int(e['sage'])])
            r = ws.max_row
            link_cell(ws.cell(r, 3), e['link'])
            ws.cell(r, 6).alignment = Alignment(horizontal='center')

    # 3. Instructions
    ws = wb.create_sheet('הוראות')
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions['A'].width = 120
    ws.column_dimensions['B'].width = 96
    fit_to_page(ws)
    for text, link, style in instructions(doc, views_note_text(views_date, bool(views))):
        ws.append([text, link])
        c = ws.cell(ws.max_row, 1)
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right', readingOrder=2)
        if style == 'title':
            c.font = Font(bold=True, size=16)
        elif style == 'head':
            c.font = Font(bold=True, size=12, color='5B4B8A')
        elif style == 'note':
            c.font = Font(italic=True, color='666666')
        if link:
            link_cell(ws.cell(ws.max_row, 2), link)

    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb.save(OUT_XLSX)


# ---- --check ------------------------------------------------------------------------------------------------
def check(doc):
    """Problems found comparing the committed files with a rebuild (views and the date they were read are ignored)."""
    from openpyxl import load_workbook
    problems = []
    if not os.path.exists(OUT_JSON) or not os.path.exists(OUT_XLSX):
        return ['%s or %s is missing: run scripts/build_description_links.py' % (OUT_JSON, OUT_XLSX)]
    old = json.load(open(OUT_JSON, encoding='utf-8'))
    for key, ident in (('site', None), ('youtube', 'row'), ('spotify', 'sage'), ('unresolved', 'sage')):
        if old.get(key) == doc[key]:
            continue
        if ident is None or not isinstance(old.get(key), list):
            problems.append('description-links.json: "%s" differs from a rebuild' % key)
            continue
        have = {e.get(ident): e for e in old[key]}
        for e in doc[key]:
            if have.get(e[ident]) != e:
                problems.append('description-links.json: %s %s differs from a rebuild (it says: %s)' % (
                    key, e[ident], (have.get(e[ident]) or {}).get('line', 'nothing')))
        for k in sorted(set(have) - {e[ident] for e in doc[key]}, key=str):
            problems.append('description-links.json: %s %s should not be there' % (key, k))
    wb = load_workbook(OUT_XLSX)
    if wb.sheetnames != ['יוטיוב', 'ספוטיפיי', 'הוראות']:
        return problems + ['the xlsx has the sheets %s' % wb.sheetnames]
    # YouTube sheet
    ws = wb['יוטיוב']
    if [c.value for c in ws[1]] != YT_HEAD:
        problems.append('יוטיוב: the header differs')
    want = {e['row']: yt_cells(e) for e in doc['youtube']}
    got, shown = {}, []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[8] is None:
            continue
        shown.append((r[1], r[8]))
        got[r[8]] = {i: r[i - 1] for i in range(3, 10)}
    if len(shown) != len(doc['youtube']):
        problems.append('יוטיוב: %d rows, a rebuild has %d' % (len(shown), len(doc['youtube'])))
    for k, cells in want.items():
        if got.get(k) != cells:
            problems.append('יוטיוב: row %s differs from a rebuild' % k)
    views = [v for v, _ in shown]
    if all(v is not None for v in views):
        if shown != sorted(shown, key=lambda t: (-t[0], t[1])):
            problems.append('יוטיוב: not sorted by views, most viewed first')
    elif any(v is not None for v in views):
        # unseen videos go last, in playlist order
        seen = [v for v in views if v is not None]
        if views[:len(seen)] != seen or seen != sorted(seen, reverse=True):
            problems.append('יוטיוב: not sorted by views, most viewed first')
    elif [r for _, r in shown] != sorted(r for _, r in shown):
        problems.append('יוטיוב: no views, and not in playlist order')
    # Spotify sheet
    ws = wb['ספוטיפיי']
    if [c.value for c in ws[1]] != SP_HEAD:
        problems.append('ספוטיפיי: the header differs')
    want = {int(e['sage']): sp_cells(e) for e in doc['spotify']}
    got, tail, total_of = {}, [], {}
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for i, r in enumerate(rows):
        if r[1] == UNRESOLVED_TITLE:
            tail = rows[i + 2:]
            rows = rows[:i]
            break
    for r in rows:
        if r[5] is not None:
            got[r[5]] = {2: r[1], 3: r[2], 4: r[3], 6: r[5]}
            total_of[r[5]] = r[4]
    if len(got) != len(doc['spotify']):
        problems.append('ספוטיפיי: %d rows, a rebuild has %d' % (len(got), len(doc['spotify'])))
    for k, cells in want.items():
        if got.get(k) != cells:
            problems.append('ספוטיפיי: sage %s differs from a rebuild' % k)
    order = [r[5] for r in rows if r[5] is not None]
    totals = [total_of[i] for i in order]
    if all(t is not None for t in totals):
        if list(zip(totals, order)) != sorted(zip(totals, order), key=lambda t: (-t[0], t[1])):
            problems.append('ספוטיפיי: not sorted by the sage\'s YouTube views, most viewed first')
    elif any(t is not None for t in totals) or order != sorted(order):
        problems.append('ספוטיפיי: no views, and not in id order')
    block = [(r[1], r[2], r[5]) for r in tail if r[5] is not None]
    if block != [(e['label'], e['link'], int(e['sage'])) for e in doc['unresolved']]:
        problems.append('ספוטיפיי: the block "%s" differs from a rebuild' % UNRESOLVED_TITLE)
    # instructions (the line that dates the views is not compared)
    ws = wb['הוראות']
    got = [(r[0] or '', r[1] or '') for r in ws.iter_rows(values_only=True)]
    want = [(t, link or '') for t, link, _ in instructions(doc, '')]
    if len(got) != len(want):
        problems.append('הוראות: %d lines, a rebuild has %d' % (len(got), len(want)))
    else:
        for i, (a, b) in enumerate(zip(got, want)):
            if i != VIEWS_NOTE_AT and a != b:
                problems.append('הוראות: line %d differs from a rebuild' % (i + 1))
    return problems


def main():
    site = read_site()
    doc = build_doc(site)
    yt, sp = doc['youtube'], doc['spotify']
    print('%d YouTube rows (he %d, en %d, ru %d) | %d Spotify episodes (+%d links that are not episode addresses) | site %s' % (
        len(yt), *(sum(1 for e in yt if e['lang'] == l) for l in LANGS), len(sp), len(doc['unresolved']), site))
    n = lambda key, rows=yt: sum(1 for e in rows if e['claims'][key])
    print('YouTube lines say: full research %d (else "page" %d), quiz %d, teachers %d, students %d | women %d, books or topics %d' % (
        n('research'), len(yt) - n('research'), n('quiz'), n('teachers'), n('students'),
        sum(1 for e in yt if e['claims']['kind'] == 'f'), sum(1 for e in yt if e['claims']['kind'] == 'n')))
    if args.check:
        problems = check(doc)
        for p in problems[:40]:
            print('  ' + p)
        if len(problems) > 40:
            print('  ... and %d more' % (len(problems) - 40))
        print('description-links: %s' % ('%d problem(s): run scripts/build_description_links.py' % len(problems) if problems else 'clean'))
        sys.exit(1 if problems else 0)

    views, views_date = None, ''
    if not args.no_stats:
        got, extra = fetch_views(doc)
        if got is None:
            print('views unavailable (%s): the YouTube sheet is in playlist order' % extra)
        else:
            views, views_date = got, datetime.date.today().isoformat()
            for w in extra:
                print('  warning: ' + w)
            seen = sum(1 for v in views.values() if v is not None)
            print('views for %d of %d videos' % (seen, len(yt)))
    os.makedirs(args.out, exist_ok=True)
    with open(OUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
        f.write('\n')
    write_xlsx(doc, views, views_date)
    print('wrote %s and %s' % (os.path.relpath(OUT_XLSX, REPO) if args.out.startswith(REPO) else OUT_XLSX,
                               os.path.relpath(OUT_JSON, REPO) if args.out.startswith(REPO) else OUT_JSON))
    if views:
        top = sorted(((v, e) for e in yt for v in [views.get(e['id'])] if v is not None), key=lambda t: -t[0])[:10]
        print('top 10 by views:')
        for v, e in top:
            print('  %6d  %s  %s  %s' % (v, e['lang'], e['id'], e['name']))


main()
